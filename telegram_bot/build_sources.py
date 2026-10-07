"""
Проверка источников из таблицы «Медиа РФ» и сборка sources.json для парсера.

Для каждого источника выясняет, КАК его читать:
  tg     — Телеграм-канал с открытой веб-лентой (t.me/s/…)
  rss    — у сайта есть RSS/Atom-лента
  html   — ленты нет, но заголовки новостей есть прямо на странице
  gnews  — сайт недоступен напрямую, но его статьи находит Google Новости (site:домен)
…или почему его читать нельзя (канала/домена не существует, закрытый чат, сайт закрыт и т.д.).

Запуск (локально, нужен xlrd для .xls):
    python telegram_bot/build_sources.py "Медиа РФ (1).xls"
Результат:
    telegram_bot/sources.json        — рабочий список для парсера
    Проверка источников.xlsx         — отчёт по каждой строке таблицы (для помощницы)
"""
import asyncio
import json
import re
import socket
import sys
import warnings
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from urllib.parse import quote, urljoin, urlparse

import httpx
from bs4 import BeautifulSoup

sys.path.insert(0, str(Path(__file__).parent))
from regions import normalize_region

warnings.filterwarnings('ignore')

HERE = Path(__file__).parent
UA = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
      '(KHTML, like Gecko) Chrome/124.0 Safari/537.36')
HEADERS = {'User-Agent': UA, 'Accept-Language': 'ru-RU,ru;q=0.9'}

REALTY_DB = ('cian.ru', 'domclick', 'etagi', 'наш.дом.рф', 'xn--80az8a.xn--d1aqf', 'erzrf', 'realty.yandex',
             '2gis', 'm2.ru', 'novostroy', 'avito', 'novostroev', 'restate', 'n1.ru', 'move.ru')
TENDERS = ('tender', 'zakupki', 'podryad', 'rostender', 'b2b-center')
FEED_PROBES = ['/rss', '/rss/', '/rss.xml', '/feed/', '/feed', '/news/rss', '/rss/news', '/yandex.xml',
               '/rss/yandex', '/news/rss.xml', '/index.xml']

# Источник «строительный» — у таких берём любую новость с объектом, у общих СМИ нужен ещё глагол стройки
CONSTRUCTION_RE = re.compile(r'стро|архит|застрой|капремонт|капитальн|градостро|\bжк\b|недвиж|девелоп|'
                             r'реновац|жилищ|госстройнадзор|ипотек|проектир|\bсро\b|инфраструкт', re.I)
CONSTRUCTION_URL_RE = re.compile(r'stro[iy]|build|nedvizh|realty|zhk|kapremont|fkr|gradostro|arhitekt', re.I)


# ── Чтение таблицы ────────────────────────────────────────────────────────────
def read_media_table(path: Path) -> list:
    import xlrd
    book = xlrd.open_workbook(str(path))
    rows = []
    for sh in book.sheets():
        hdr = [str(h).strip() for h in sh.row_values(0)]
        if 'Ссылка' not in hdr:
            continue                       # справочные листы со списком субъектов
        region = None
        for r in range(1, sh.nrows):
            d = dict(zip(hdr, [str(c).strip() for c in sh.row_values(r)]))
            if sh.name == 'Миллионники':
                region = d.get('Название города') or region
                name, kind = d.get('Название', ''), d.get('Тип', '')
            else:
                region = d.get('Область') or region
                name, kind = d.get('Название портала/канала', ''), ''
            url = d.get('Ссылка', '')
            if not url and not name:
                continue
            rows.append({'sheet': sh.name, 'row': r + 1, 'region': normalize_region(region),
                         'name': name, 'kind': kind, 'topic': d.get('Тематика', ''), 'url': url})
    return rows


def classify_url(u: str) -> str:
    ul = u.lower()
    if not ul.startswith('http'):
        return 'no_url'
    if 't.me/' in ul:
        return 'telegram'
    if 'vk.com' in ul or 'vk.ru' in ul:
        return 'vk'
    if 'ok.ru' in ul or 'youtube' in ul or 'rutube' in ul or 'max.ru' in ul:
        return 'other_social'
    if 'dzen.ru' in ul:
        return 'dzen'
    if 'tgstat' in ul:
        return 'catalog'
    if any(x in ul for x in TENDERS):
        return 'tenders'
    if any(x in ul for x in REALTY_DB):
        return 'realty_db'
    return 'website'


def tg_handle(u: str):
    m = re.search(r't\.me/(?:s/)?([A-Za-z0-9_]+)', u)
    return m.group(1) if m and m.group(1).lower() not in ('joinchat', 'addlist') else None


def host_of(u: str) -> str:
    h = (urlparse(u).hostname or '').lower()
    return h[4:] if h.startswith('www.') else h


# ── Проверка Телеграма ────────────────────────────────────────────────────────
async def check_tg(client, handle):
    r = await client.get(f'https://t.me/s/{handle}')
    soup = BeautifulSoup(r.text, 'html.parser')
    if soup.find('div', class_='tgme_widget_message_wrap'):
        dates = [t.get('datetime', '')[:10] for t in soup.find_all('time') if t.get('datetime')]
        return 'tg', f'последний пост {max(dates) if dates else "?"}'
    page = (await client.get(f'https://t.me/{handle}')).text
    title = re.search(r'tgme_page_title[^>]*>\s*<span[^>]*>([^<]*)', page)
    extra = re.search(r'tgme_page_extra[^>]*>([^<]*)', page)
    if not title:
        return 'fail', 'канала с таким адресом не существует'
    extra = extra.group(1) if extra else ''
    if 'member' in extra:
        return 'fail', f'это чат-группа ({extra.strip()}), не канал — без аккаунта не читается'
    if 'subscriber' in extra:
        return 'fail', f'у канала отключён веб-просмотр ({extra.strip()})'
    return 'fail', 'это личный аккаунт или бот, а не канал'


# ── Проверка сайтов ───────────────────────────────────────────────────────────
def _is_feed(text: str) -> bool:
    head = text[:3000].lower()
    return '<rss' in head or '<feed' in head or '<rdf' in head


def _feed_items(text: str) -> int:
    return len(re.findall(r'<(?:item|entry)[\s>]', text))


async def _fetch(client, url):
    try:
        return await client.get(url)
    except Exception as e:
        return e


async def check_site(client, url):
    r = await _fetch(client, url)
    if isinstance(r, Exception):
        return 'unreachable', type(r).__name__
    if r.status_code >= 400:
        return 'unreachable', f'HTTP {r.status_code}'
    final = str(r.url)
    if _is_feed(r.text) and _feed_items(r.text):
        return 'rss', final
    soup = BeautifulSoup(r.text, 'lxml')
    cands = [urljoin(final, l['href']) for l in soup.find_all('link', type=re.compile('rss|atom', re.I))
             if l.get('href')]
    cands += [urljoin(final, a['href']) for a in soup.find_all('a', href=re.compile(r'rss|feed', re.I))]
    root = f'{urlparse(final).scheme}://{urlparse(final).netloc}'
    cands += [root + p for p in FEED_PROBES]
    seen = set()
    for fu in cands:
        if fu in seen or 'vk.com' in fu or 't.me' in fu:
            continue
        seen.add(fu)
        if len(seen) > 16:
            break
        fr = await _fetch(client, fu)
        if not isinstance(fr, Exception) and fr.status_code == 200 and _is_feed(fr.text) and _feed_items(fr.text):
            return 'rss', fu
    host = urlparse(final).netloc
    heads = {a.get_text(' ', strip=True) for a in soup.find_all('a', href=True)
             if 30 <= len(a.get_text(' ', strip=True)) <= 220
             and urlparse(urljoin(final, a['href'])).netloc in (host, '')}
    if len(heads) >= 5:
        return 'html', final
    return 'no_content', 'страница пустая (содержимое грузится скриптом) или новостей нет'


async def check_gnews(client, host):
    q = quote(f'site:{host} when:30d')
    try:
        r = await client.get(f'https://news.google.com/rss/search?q={q}&hl=ru&gl=RU&ceid=RU:ru')
        return r.text.count('<item>')
    except Exception:
        return 0


def dns_ok(host: str) -> bool:
    try:
        socket.getaddrinfo(host.encode('idna').decode(), 443)
        return True
    except Exception:
        return False


# ── Главная процедура ─────────────────────────────────────────────────────────
async def run_checks(rows: list):
    for r in rows:
        r['cls'] = classify_url(r['url'])

    handles = sorted({tg_handle(r['url']) for r in rows if r['cls'] == 'telegram'} - {None})
    sites = sorted({r['url'] for r in rows if r['cls'] in ('website', 'dzen')})
    print(f'Телеграм-каналов: {len(handles)}, сайтов: {len(sites)}')

    tg_res, site_res, gn_res = {}, {}, {}
    sem_tg, sem_web, sem_gn = asyncio.Semaphore(6), asyncio.Semaphore(40), asyncio.Semaphore(4)
    limits = httpx.Limits(max_connections=60)
    async with httpx.AsyncClient(headers=HEADERS, timeout=httpx.Timeout(20, connect=12),
                                 follow_redirects=True, verify=False, limits=limits) as client:
        async def one_tg(h):
            async with sem_tg:
                try:
                    tg_res[h] = await check_tg(client, h)
                except Exception as e:
                    tg_res[h] = ('fail', f'ошибка проверки: {type(e).__name__}')

        async def one_site(u):
            async with sem_web:
                site_res[u] = await check_site(client, u)

        await asyncio.gather(*[one_tg(h) for h in handles], *[one_site(u) for u in sites])
        print('  прямые проверки готовы, проверяю недоступные сайты через DNS и Google Новости…')

        bad_hosts = sorted({host_of(u) for u, (st, _) in site_res.items() if st in ('unreachable', 'no_content')})
        with ThreadPoolExecutor(30) as ex:
            dns = dict(zip(bad_hosts, ex.map(dns_ok, bad_hosts)))

        async def one_gn(h):
            async with sem_gn:
                gn_res[h] = await check_gnews(client, h)
                await asyncio.sleep(0.5)

        await asyncio.gather(*[one_gn(h) for h in bad_hosts if dns[h]])

    # Итог по каждой строке таблицы
    for r in rows:
        c = r['cls']
        if c == 'telegram':
            h = tg_handle(r['url'])
            st, note = tg_res.get(h, ('fail', 'приватная ссылка-приглашение — без аккаунта не читается'))
            r.update(method='tg' if st == 'tg' else '', target=h, note=note)
        elif c in ('website', 'dzen'):
            st, info = site_res[r['url']]
            if st in ('rss', 'html'):
                r.update(method=st, target=info, note='RSS-лента' if st == 'rss' else 'заголовки со страницы')
            else:
                h = host_of(r['url'])
                if not dns.get(h, True):
                    r.update(method='', target='', note='такого сайта не существует (домен не найден)')
                elif gn_res.get(h, 0) > 0:
                    r.update(method='gnews', target=h,
                             note=f'напрямую не открывается ({info}) — читаем через Google Новости')
                elif info.startswith('HTTP 404'):
                    r.update(method='', target='', note='страница не найдена (404) — проверьте ссылку')
                elif st == 'no_content':
                    r.update(method='', target='', note=info)
                else:
                    r.update(method='', target='',
                             note=f'сайт закрыт для зарубежных серверов ({info}) и не индексируется Google Новостями')
        else:
            r.update(method='', target='', note={
                'vk': 'ВКонтакте — нужен ключ доступа VK API (подключим отдельно)',
                'realty_db': 'база новостроек/объявлений, а не лента новостей — в парсер не подходит',
                'tenders': 'тендерная площадка — в парсер новостей не подходит',
                'catalog': 'каталог каналов, а не источник новостей',
                'other_social': 'Одноклассники/YouTube/MAX — не поддерживается',
                'no_url': 'нет ссылки',
            }.get(c, ''))
    return rows


def write_sources(rows: list, path: Path):
    """sources.json: по одному источнику на (метод, адрес), регион — первый встретившийся."""
    out, seen = [], set()
    for r in rows:
        if not r['method']:
            continue
        key = (r['method'], r['target'])
        if key in seen:
            continue
        seen.add(key)
        out.append({
            'method': r['method'], 'target': r['target'], 'region': r['region'],
            'name': r['name'] or r['target'],
            # Тематику не учитываем: у помощницы почти в каждой строке написано «строительство»,
            # хотя сама лента общая. Смотрим на название, тип и адрес самого источника.
            'construction': bool(CONSTRUCTION_RE.search(f"{r['name']} {r['kind']}")
                                 or CONSTRUCTION_URL_RE.search(f"{r['url']} {r['target']}")),
        })
    path.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding='utf-8')
    print(f'\nsources.json: {len(out)} источников — ' +
          ', '.join(f'{k}: {v}' for k, v in Counter(s["method"] for s in out).most_common()))


def write_report(rows: list, path: Path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Все источники'
    cols = ['Лист', 'Строка', 'Регион', 'Название', 'Ссылка', 'Подключён', 'Как читаем', 'Комментарий']
    ws.append(cols)
    labels = {'tg': 'Телеграм', 'rss': 'RSS', 'html': 'Страница сайта', 'gnews': 'Google Новости'}
    for r in rows:
        ws.append([r['sheet'], r['row'], r['region'], r['name'], r['url'],
                   'да' if r['method'] else 'нет', labels.get(r['method'], ''), r['note']])
    green, red = PatternFill('solid', fgColor='E2EFDA'), PatternFill('solid', fgColor='FCE4D6')
    for row in ws.iter_rows(min_row=2):
        row[5].fill = green if row[5].value == 'да' else red
    for c, w in zip('ABCDEFGH', (14, 8, 26, 40, 45, 11, 16, 70)):
        ws.column_dimensions[c].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F4E78')
        c.alignment = Alignment(horizontal='center')
    t = Table(displayName='Sources', ref=f'A1:H{ws.max_row}')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = 'A2'

    s = wb.create_sheet('Итог')
    s.append(['Результат', 'Строк'])
    for k, v in Counter(r['note'] if not r['method'] else f'Подключён: {labels[r["method"]]}'
                        for r in rows).most_common():
        s.append([k, v])
    s.column_dimensions['A'].width = 95
    s.append([])
    s.append([f'Проверено {datetime.now():%d.%m.%Y %H:%M}, всего строк: {len(rows)}'])
    wb.move_sheet('Итог', offset=-1)
    wb.save(path)
    print(f'Отчёт: {path}')


def main():
    src = Path(sys.argv[1] if len(sys.argv) > 1 else 'Медиа РФ (1).xls')
    rows = read_media_table(src)
    print(f'Строк в таблице: {len(rows)}')
    rows = asyncio.run(run_checks(rows))
    write_sources(rows, HERE / 'sources.json')
    write_report(rows, HERE.parent / 'Проверка источников.xlsx')


if __name__ == '__main__':
    main()
