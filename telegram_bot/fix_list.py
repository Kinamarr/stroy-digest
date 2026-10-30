"""
Список источников «на исправление» для помощницы и подстановка её исправлений.

1) Сделать список (из отчёта build_sources.py):
       python telegram_bot/fix_list.py make
   → «Исправить источники.xlsx»: только строки, которые можно починить, с подсказкой и пустой колонкой
     «Новая ссылка».
2) Когда помощница заполнит колонку «Новая ссылка», подставить её в таблицу источников:
       python telegram_bot/fix_list.py apply "Медиа РФ (1).xls"
   → «Медиа РФ (исправлено).xlsx», которую затем прогнать через build_sources.py.
"""
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).parent.parent
REPORT = ROOT / 'Проверка источников.xlsx'
FIX = ROOT / 'Исправить источники.xlsx'

# Комментарий из отчёта → что сделать помощнице
HINTS = [
    ('канала с таким адресом не существует',
     'Найти настоящий канал в Телеграме (поиск по названию) и вставить ссылку вида https://t.me/имя'),
    ('личный аккаунт или бот',
     'Это не канал. Найти официальный канал организации в Телеграме'),
    ('такого сайта не существует',
     'Найти настоящий сайт (поиск по названию). Лучше сразу раздел новостей: …/news/'),
    ('страница не найдена (404)',
     'Открыть сайт и найти действующий раздел новостей, вставить его адрес'),
    ('страница пустая',
     'Найти у этого СМИ/организации Телеграм-канал или страницу «Новости» и вставить её'),
    ('группы ВКонтакте с таким адресом не существует',
     'Найти настоящую группу ВКонтакте и вставить ссылку вида https://vk.com/имя'),
    ('закрытая группа ВКонтакте',
     'Найти открытую группу/канал этой организации (ВК или Телеграм)'),
    ('это чат-группа',
     'Это чат, а не канал. Найти новостной канал этой организации'),
    ('у канала отключён веб-просмотр',
     'Найти другой канал этой организации или её сайт'),
    ('нет ссылки', 'Найти и вставить ссылку'),
    ('сайт закрыт для зарубежных серверов',
     'Если у организации есть Телеграм-канал или группа ВК — вставить ссылку на них'),
]


def make():
    ws = openpyxl.load_workbook(REPORT)['Все источники']
    hdr = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, r))
        if d['Подключён'] == 'да':
            continue
        hint = next((h for key, h in HINTS if key in (d['Комментарий'] or '')), None)
        if hint:
            rows.append([d['Регион'], d['Название'], d['Ссылка'], d['Комментарий'], hint, None,
                         d['Лист'], d['Строка']])
    rows.sort(key=lambda r: (r[0] or '', r[1] or ''))

    wb = openpyxl.Workbook()
    out = wb.active
    out.title = 'Исправить'
    cols = ['Регион', 'Название', 'Текущая ссылка', 'Что не так', 'Что сделать', 'Новая ссылка',
            'Лист', 'Строка']
    out.append(cols)
    for r in rows:
        out.append(r)
    widths = (24, 40, 40, 45, 55, 40, 12, 8)
    for i, w in enumerate(widths):
        out.column_dimensions['ABCDEFGH'[i]].width = w
    for c in out[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F4E78')
        c.alignment = Alignment(horizontal='center', wrap_text=True)
    yellow = PatternFill('solid', fgColor='FFF2CC')
    for row in out.iter_rows(min_row=2):
        row[5].fill = yellow
        for c in row[:5]:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    t = Table(displayName='Fix', ref=f'A1:H{out.max_row}')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
    out.add_table(t)
    out.freeze_panes = 'C2'
    wb.save(FIX)
    print(f'{FIX.name}: {len(rows)} строк на исправление')


def apply(media_path: Path):
    """Подставляет «Новую ссылку» в копию таблицы источников (лист/строка — из отчёта)."""
    import xlrd
    fixes = {}
    ws = openpyxl.load_workbook(FIX)['Исправить']
    hdr = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, r))
        if d['Новая ссылка']:
            fixes[(d['Лист'], int(d['Строка']))] = str(d['Новая ссылка']).strip()
    book = xlrd.open_workbook(str(media_path))
    out = openpyxl.Workbook()
    out.remove(out.active)
    done = 0
    for sh in book.sheets():
        o = out.create_sheet(sh.name[:31])
        hdr = [str(h).strip() for h in sh.row_values(0)]
        link_col = hdr.index('Ссылка') if 'Ссылка' in hdr else None
        for r in range(sh.nrows):
            vals = sh.row_values(r)
            if link_col is not None and (sh.name, r + 1) in fixes:
                vals[link_col] = fixes[(sh.name, r + 1)]
                done += 1
            o.append(vals)
    dst = media_path.with_name('Медиа РФ (исправлено).xlsx')
    out.save(dst)
    print(f'Исправлено ссылок: {done} → {dst.name}. Дальше: python telegram_bot/build_sources.py "{dst.name}"')


if __name__ == '__main__':
    if sys.argv[1:2] == ['make']:
        make()
    elif sys.argv[1:2] == ['apply']:
        apply(Path(sys.argv[2]))
    else:
        print(__doc__)
